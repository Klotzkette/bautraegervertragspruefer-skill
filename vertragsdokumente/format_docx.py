#!/usr/bin/env python3
"""Apply the shared, editable draft layout after a contract conversion."""
from __future__ import annotations

import sys
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def _keep_leading_rows_with_content(table) -> None:
    """Put a standalone heading/signature rule and its text in the same row.

    Paragraph keep-with-next is not reliable across table-row boundaries in
    LibreOffice. Move the existing OOXML blocks column by column instead of
    using cell.merge(), which can interleave the two language versions.
    """
    def leading_cell(cell) -> bool:
        paragraphs = [p for p in cell.paragraphs if p.text.strip()]
        if len(paragraphs) != 1 or cell.tables:
            return False
        paragraph = paragraphs[0]
        text = paragraph.text.strip()
        if re.fullmatch(r"[_\-–—\s]{3,}", text):
            return True
        if paragraph.style.name.startswith("Heading"):
            return True
        # Numbered subheadings were converted to Table Contents, not Heading.
        return bool(
            re.fullmatch(r"\d+(?:\.\d+)+\s+[^.!?;:\n]{1,100}", text)
            and len(text.split()) <= 8
        )

    index = 1  # The first row is the repeating language header.
    while index + 1 < len(table.rows):
        row, following = table.rows[index], table.rows[index + 1]
        cells, next_cells = row._tr.tc_lst, following._tr.tc_lst
        ordinary_pair = (
            len(cells) == len(next_cells) == 2
            and not any(tc.xpath("./w:tcPr/w:gridSpan | ./w:tcPr/w:vMerge")
                        for tc in [*cells, *next_cells])
        )
        if ordinary_pair and all(leading_cell(cell) for cell in row.cells):
            for cell, next_cell in zip(row.cells, following.cells):
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
                for block in list(next_cell._tc):
                    if block.tag != qn("w:tcPr"):
                        cell._tc.append(block)
            following._tr.getparent().remove(following._tr)
        index += 1


def format_contract(path: Path) -> None:
    document = Document(path)
    bilingual = path.stem.endswith("-de-en")
    if "Title" not in document.styles:
        title = document.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
        title.font.name = "Times New Roman"
        title.font.size = Pt(16)
    for style in document.styles:
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
            for color in style.element.xpath(".//w:color"):
                for attribute in ("themeColor", "themeTint", "themeShade"):
                    color.attrib.pop(qn(f"w:{attribute}"), None)
            for border in style.element.xpath(".//w:pBdr"):
                border.getparent().remove(border)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10 if bilingual else 11)
    normal.paragraph_format.line_spacing = 1.12
    for section in document.sections:
        if not bilingual:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.2)
            section.left_margin, section.right_margin = Cm(2.5), Cm(2.2)
        else:
            section.page_width, section.page_height = Cm(29.7), Cm(21)
            section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.5)
            section.left_margin, section.right_margin = Cm(1.4), Cm(1.4)
            section.header_distance, section.footer_distance = Cm(0.6), Cm(0.6)
        header = section.header.paragraphs[0]
        header.text = "ENTWURF / DRAFT" if bilingual else "ENTWURF"
        header.runs[0].bold = True
        header.runs[0].font.size = Pt(9)
    for paragraph in document.paragraphs:
        if paragraph.text == "Bauträgervertrag" or "Deutsch/English" in paragraph.text:
            paragraph.style = document.styles["Title"]
        for border in paragraph._p.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    for node in document.element.xpath(".//w:color"):
        node.set(qn("w:val"), "000000")
        for attribute in ("themeColor", "themeTint", "themeShade"):
            node.attrib.pop(qn(f"w:{attribute}"), None)
    for table in document.tables:
        if bilingual:
            _keep_leading_rows_with_content(table)
            for row in table.rows:
                properties = row._tr.get_or_add_trPr()
                if properties.find(qn("w:cantSplit")) is None:
                    properties.append(OxmlElement("w:cantSplit"))
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.widow_control = True
                        if re.fullmatch(r"[_\-–—\s]{3,}", paragraph.text):
                            paragraph.paragraph_format.keep_with_next = True
        if not bilingual and len(table.columns) == 3 and "Rate" in table.rows[0].cells[0].text:
            table.autofit = False
            widths = (Cm(1.3), Cm(12.4), Cm(2.6))
            for column, width in zip(table.columns, widths):
                column.width = width
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = width
            row_properties = table.rows[0]._tr.get_or_add_trPr()
            if row_properties.find(qn("w:tblHeader")) is None:
                row_properties.append(OxmlElement("w:tblHeader"))
        properties = table._tbl.tblPr
        borders = properties.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            properties.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = borders.find(qn(f"w:{edge}"))
            if border is None:
                border = OxmlElement(f"w:{edge}")
                borders.append(border)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                border.set(qn(f"w:{key}"), value)
    document.core_properties.title = "Bauträgervertrag Entwurf"
    document.save(path)


if __name__ == "__main__":
    format_contract(Path(sys.argv[1]))
