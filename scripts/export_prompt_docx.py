#!/usr/bin/env python3
"""Export a complete portable prompt to readable Word, with a real page gate.

Requires python-docx, pandoc, LibreOffice and pdfinfo. Pass --soffice explicitly
when a managed runtime is available. No source file is rewritten. PDF output is
temporary QA material, not a second deliverable. Markdown has no intrinsic page
count; --max-pages checks this export in the chosen LibreOffice installation.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile


def frontmatter_body(text: str) -> str:
    """Omit plugin metadata, but retain the complete human-facing prompt."""
    if text.startswith("---\n"):
        match = re.search(r"\n---\s*\n", text[4:])
        if match:
            return text[4 + match.end():].lstrip("\n")
    return text


def narrative_tables(text: str) -> str:
    """Lay out long narrative records vertically without shortening their text."""
    lines = text.splitlines()
    out = []
    pos = 0
    fence = None
    while pos < len(lines):
        line = lines[pos]
        if fence is not None:
            # A shorter fence or a different marker is literal code, not a close.
            marker, length = fence
            if re.fullmatch(r" {0,3}" + re.escape(marker) + "{" + str(length) + r",}[ \t]*", line):
                fence = None
            out.append(line)
            pos += 1
            continue
        opening = re.match(r"^ {0,3}(`{3,}|~{3,})(.*)$", line)
        if opening and not (opening[1][0] == "`" and "`" in opening[2]):
            fence = (opening[1][0], len(opening[1]))
            out.append(line)
            pos += 1
            continue
        if not lines[pos].startswith("|"):
            out.append(lines[pos])
            pos += 1
            continue
        end = pos
        while end < len(lines) and lines[end].startswith("|"):
            end += 1
        group = lines[pos:end]
        rows = [[cell.strip() for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
                for line in group]
        headers = rows[0]
        if (len(rows) >= 3 and len(headers) >= 3
                and all(len(row) == len(headers) for row in rows)
                and all(re.fullmatch(r":?-+:?", cell) for cell in rows[1])
                and max(len(cell) for row in rows[2:] for cell in row) > 400):
            for row in rows[2:]:
                out.extend(["", "### " + row[0], ""])
                for label, value in zip(headers[1:], row[1:]):
                    out.extend(["**" + label + ":** " + value, ""])
        else:
            out.extend(group)
        pos = end
    return "\n".join(out) + "\n"


def format_docx(path: Path, paper: str = "Letter") -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    section = doc.sections[0]
    if paper == "A4":
        section.page_width, section.page_height = Cm(21), Cm(29.7)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    content_width = (section.page_width - section.left_margin - section.right_margin) / 360000
    section.header_distance = section.footer_distance = Cm(1)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name in ("Body Text", "First Paragraph", "Compact"):
        if name in doc.styles:
            doc.styles[name].base_style = normal
            doc.styles[name].font.name = "Arial"
            doc.styles[name].font.size = Pt(11)
            doc.styles[name].paragraph_format.space_after = Pt(5)
            doc.styles[name].paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 19), ("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Title"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        for border in list(style.element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    if doc.paragraphs:
        doc.paragraphs[0].style = doc.styles["Title"]
    for paragraph in doc.paragraphs:
        if paragraph.style.name not in {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"}:
            paragraph.paragraph_format.keep_with_next = False
        for run in paragraph.runs:
            # Inline Markdown code is an instruction phrase, not a program listing.
            run.font.name = "Arial"
    for table in doc.tables:
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            caption = Paragraph(previous, doc)
            if 0 < len(caption.text) < 180:
                caption.paragraph_format.keep_with_next = True
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        count = len(table.columns)
        if count == 2:
            fractions = [.3, .7]
        elif count == 3:
            fractions = [.48, .20, .32]
        else:
            fractions = [1 / count] * count
        for col, fraction in zip(table.columns, fractions):
            col.width = Cm(content_width * fraction)
        props = table._tbl.tblPr
        borders = props.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            props.append(borders)
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            item = OxmlElement("w:" + side)
            for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                item.set(qn("w:" + key), value)
            borders.append(item)
        for rowno, row in enumerate(table.rows):
            # A one-word cell continuation on the next page is not useful.
            # Tables may span pages; individual short rows stay intact.
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            if rowno == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            for colno, cell in enumerate(row.cells):
                cell.width = Cm(content_width * fractions[colno])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                for existing in list(tcpr.findall(qn("w:tcBorders"))):
                    tcpr.remove(existing)
                cell_borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + side)
                    for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                        border.set(qn("w:" + key), value)
                    cell_borders.append(border)
                tcpr.append(cell_borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    margin = OxmlElement("w:" + side)
                    margin.set(qn("w:w"), "80")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                tcpr.append(margins)
                if rowno == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E7E6E6")
                    tcpr.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_with_next = rowno == 0
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Arial", Pt(11)
                        if rowno == 0:
                            run.font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(path)


def document_text(path: Path) -> tuple[str, ...]:
    """Ordered text nodes include paragraph and table contents, but not metadata."""
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return tuple(node.text or "" for node in root.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--max-pages", type=int, default=100)
    parser.add_argument("--paper", choices=["Letter", "A4"], default="Letter")
    parser.add_argument("--soffice", default=os.environ.get("BTV_SOFFICE", "soffice"))
    parser.add_argument("--keep-wide-tables", action="store_true", help="Diagnostic baseline only")
    args = parser.parse_args()
    if args.max_pages <= 0 or args.output.suffix.lower() != ".docx":
        parser.error("A positive page limit and a .docx output are required")
    if args.source.resolve() == args.output.resolve():
        parser.error("Source must not be overwritten")
    if args.output.exists():
        parser.error("Output exists; choose a new path to preserve it")
    source_bytes = args.source.read_bytes()
    content = source_bytes.decode("utf-8")
    body = frontmatter_body(content)
    if not args.keep_wide_tables:
        body = narrative_tables(body)
    # Color emoji can disappear in Office PDF renderers. Explicit labels remain
    # legible, searchable and semantically identical, including in monochrome.
    for symbol, label in {"🔴": "ROT", "🟠": "ORANGE", "🟢": "GRÜN"}.items():
        body = body.replace(symbol, label)
    # Permit a compound-boundary break in this narrow table label instead of
    # stranding its last letter. The source Markdown is never modified.
    body = body.replace("| Vollstreckungsunterwerfung |",
                        "| Vollstreckungs\u00adunterwerfung |")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="btv-prompt-word-") as temp:
        work = Path(temp)
        staged = work / args.output.name
        subprocess.run(["pandoc", "--from=markdown", "--to=docx", "-o", str(staged)],
                       input=body, text=True, check=True)
        expected_text = document_text(staged)
        format_docx(staged, args.paper)
        if document_text(staged) != expected_text:
            raise RuntimeError("Word formatting changed source text; export rejected")
        # A private LO profile avoids interference with the user's open documents.
        subprocess.run([args.soffice, "-env:UserInstallation=" + (work / "profile").as_uri(),
                        "--headless", "--convert-to", "pdf", "--outdir", str(work), str(staged)],
                       check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        pdf = staged.with_suffix(".pdf")
        if not pdf.is_file():
            raise RuntimeError("LibreOffice produced no PDF; no Word export was delivered")
        info = subprocess.run(["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout
        match = re.search(r"^Pages:\s+(\d+)\s*$", info, re.MULTILINE)
        if not match:
            raise RuntimeError("Could not determine the rendered page count")
        pages = int(match.group(1))
        if pages > args.max_pages:
            raise RuntimeError(f"Rendered {pages} pages; limit is {args.max_pages}. Content must be revised.")
        shutil.copyfile(staged, args.output)
    print(json.dumps({"output": str(args.output.resolve()), "pages": pages,
                      "max_pages": args.max_pages, "font": "Arial 11 pt", "paper": args.paper,
                      "margins": "1 inch on every side",
                      "source_sha256": hashlib.sha256(source_bytes).hexdigest(),
                      "docx_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
                      "content_integrity": "all ordered main-document text nodes preserved",
                      "visual_qa": "pending; render and inspect every page before delivery"},
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
